import csv
import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib import messages
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from accounts.decorators import role_required

from .forms import AccountingSetupForm, CashTransactionForm
from .models import CashTransaction
from .services import account_balance, create_transaction, get_settings, ledger_for

ROLES = ("admin", "accountant")
MAX_UPLOAD_BYTES = 5 * 1024 * 1024


def _parse_date(raw, default=None):
    try:
        return datetime.date.fromisoformat((raw or "").strip())
    except ValueError:
        return default


def _month_start():
    today = timezone.localdate()
    return today.replace(day=1)


def _filtered_entries(request):
    school = request.school
    qs = CashTransaction.objects.filter(school=school).select_related(
        "created_by", "voided_by", "source_payment"
    )
    transaction_type = request.GET.get("type", "")
    account_type = request.GET.get("account", "")
    category = request.GET.get("category", "")
    status = request.GET.get("status", "active")
    search = request.GET.get("q", "").strip()
    date_from = _parse_date(request.GET.get("date_from"))
    date_to = _parse_date(request.GET.get("date_to"))

    if transaction_type in ("receipt", "payment"):
        qs = qs.filter(transaction_type=transaction_type)
    if account_type in ("cash", "bank"):
        qs = qs.filter(account_type=account_type)
    if category:
        qs = qs.filter(category=category)
    if status == "voided":
        qs = qs.filter(is_voided=True)
    elif status != "all":
        status = "active"
        qs = qs.filter(is_voided=False)
    if date_from:
        qs = qs.filter(entry_date__gte=date_from)
    if date_to:
        qs = qs.filter(entry_date__lte=date_to)
    if search:
        qs = qs.filter(
            Q(party__icontains=search)
            | Q(reference__icontains=search)
            | Q(description__icontains=search)
            | Q(voucher_no__icontains=search)
        )
    qs = qs.order_by("-entry_date", "-id")
    return qs, {
        "transaction_type": transaction_type,
        "account_type": account_type,
        "category": category,
        "status": status,
        "q": search,
        "date_from": date_from.isoformat() if date_from else "",
        "date_to": date_to.isoformat() if date_to else "",
    }


@role_required(*ROLES)
def dashboard(request):
    school = request.school
    settings_obj = get_settings(school)
    today = timezone.localdate()
    month_start = _month_start()
    cash_balance = account_balance(
        school, CashTransaction.ACCOUNT_CASH, settings_obj=settings_obj
    )
    bank_balance = account_balance(
        school, CashTransaction.ACCOUNT_BANK, settings_obj=settings_obj
    )
    month_qs = CashTransaction.objects.filter(
        school=school,
        is_voided=False,
        entry_date__gte=month_start,
        entry_date__lte=today,
    )
    month_totals = month_qs.aggregate(
        receipts=Sum(
            "amount",
            filter=Q(transaction_type=CashTransaction.TYPE_RECEIPT),
        ),
        payments=Sum(
            "amount",
            filter=Q(transaction_type=CashTransaction.TYPE_PAYMENT),
        ),
    )
    month_receipts = month_totals["receipts"] or Decimal("0.00")
    month_payments = month_totals["payments"] or Decimal("0.00")

    return render(request, "accounting/dashboard.html", {
        "settings_obj": settings_obj,
        "cash_balance": cash_balance,
        "bank_balance": bank_balance,
        "total_balance": cash_balance + bank_balance,
        "month_receipts": month_receipts,
        "month_payments": month_payments,
        "month_net": month_receipts - month_payments,
        "month_start": month_start,
        "recent_entries": CashTransaction.objects.filter(
            school=school, is_voided=False
        ).select_related("created_by")[:8],
        "today": today,
    })


@role_required(*ROLES)
def setup(request):
    school = request.school
    settings_obj = get_settings(school)
    form = AccountingSetupForm(
        request.POST or None,
        school=school,
        instance=settings_obj,
    )
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Opening balances and bank details updated.")
        return redirect("accounting:dashboard")
    return render(request, "accounting/setup.html", {
        "form": form,
        "settings_obj": settings_obj,
    })


@role_required(*ROLES)
def entry_list(request):
    qs, filters = _filtered_entries(request)
    totals = qs.aggregate(
        receipts=Sum(
            "amount",
            filter=Q(transaction_type=CashTransaction.TYPE_RECEIPT),
        ),
        payments=Sum(
            "amount",
            filter=Q(transaction_type=CashTransaction.TYPE_PAYMENT),
        ),
    )
    params = request.GET.copy()
    params.pop("page", None)
    page_obj = Paginator(qs, 50).get_page(request.GET.get("page"))
    total_receipts = totals["receipts"] or Decimal("0.00")
    total_payments = totals["payments"] or Decimal("0.00")
    return render(request, "accounting/entries.html", {
        **filters,
        "page_obj": page_obj,
        "categories": CashTransaction.CATEGORY_CHOICES,
        "total_receipts": total_receipts,
        "total_payments": total_payments,
        "net_total": total_receipts - total_payments,
        "querystring": params.urlencode(),
    })


@role_required(*ROLES)
def entry_create(request):
    transaction_type = request.GET.get("type", CashTransaction.TYPE_RECEIPT)
    if transaction_type not in (
        CashTransaction.TYPE_RECEIPT,
        CashTransaction.TYPE_PAYMENT,
    ):
        transaction_type = CashTransaction.TYPE_RECEIPT
    initial = {
        "transaction_type": transaction_type,
        "entry_date": timezone.localdate(),
        "account_type": CashTransaction.ACCOUNT_CASH,
        "category": (
            CashTransaction.CATEGORY_STUDENT_FEES
            if transaction_type == CashTransaction.TYPE_RECEIPT
            else "supplies"
        ),
    }
    form = CashTransactionForm(
        request.POST or None,
        school=request.school,
        initial=initial,
    )
    if request.method == "POST" and form.is_valid():
        data = form.cleaned_data
        entry = create_transaction(
            school=request.school,
            transaction_type=data["transaction_type"],
            account_type=data["account_type"],
            category=data["category"],
            party=data["party"],
            amount=data["amount"],
            entry_date=data["entry_date"],
            reference=data["reference"],
            description=data["description"],
            created_by=request.user,
        )
        messages.success(
            request,
            f"{entry.document_label} {entry.voucher_no} recorded.",
        )
        return redirect("accounting:entry_detail", pk=entry.pk)
    return render(request, "accounting/entry_form.html", {
        "form": form,
        "target": None,
    })


@role_required(*ROLES)
def entry_edit(request, pk):
    entry = get_object_or_404(
        CashTransaction.objects.filter(school=request.school, is_voided=False),
        pk=pk,
    )
    if entry.source_payment_id:
        messages.error(
            request,
            "Student fee receipts are corrected from Payment Status, not here.",
        )
        return redirect("accounting:entry_detail", pk=entry.pk)
    form = CashTransactionForm(
        request.POST or None,
        school=request.school,
        instance=entry,
    )
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, f"{entry.voucher_no} updated.")
        return redirect("accounting:entry_detail", pk=entry.pk)
    return render(request, "accounting/entry_form.html", {
        "form": form,
        "target": entry,
    })


@role_required(*ROLES)
def entry_detail(request, pk):
    entry = get_object_or_404(
        CashTransaction.objects.select_related(
            "created_by", "voided_by", "source_payment"
        ).filter(school=request.school),
        pk=pk,
    )
    return render(request, "accounting/entry_detail.html", {
        "entry": entry,
    })


@role_required(*ROLES)
@require_POST
def entry_void(request, pk):
    entry = get_object_or_404(
        CashTransaction.objects.filter(
            school=request.school, is_voided=False
        ),
        pk=pk,
    )
    if entry.source_payment_id:
        messages.error(
            request,
            "This receipt came from a student fee payment. Void it from Payment Status.",
        )
        return redirect("accounting:entry_detail", pk=entry.pk)
    entry.is_voided = True
    entry.voided_at = timezone.now()
    entry.voided_by = request.user
    entry.void_reason = request.POST.get("reason", "").strip()[:240]
    entry.save(
        update_fields=["is_voided", "voided_at", "voided_by", "void_reason"]
    )
    messages.success(request, f"{entry.voucher_no} voided and removed from the ledger.")
    return redirect("accounting:entry_detail", pk=entry.pk)


@role_required(*ROLES)
def ledger(request):
    school = request.school
    account_type = request.GET.get("account", CashTransaction.ACCOUNT_CASH)
    if account_type not in (
        CashTransaction.ACCOUNT_CASH,
        CashTransaction.ACCOUNT_BANK,
    ):
        account_type = CashTransaction.ACCOUNT_CASH
    date_from = _parse_date(request.GET.get("date_from"), _month_start())
    date_to = _parse_date(request.GET.get("date_to"), timezone.localdate())
    if date_from > date_to:
        date_from, date_to = date_to, date_from
        messages.info(request, "The date range was adjusted so the start is before the end.")

    data = ledger_for(school, account_type, date_from, date_to)
    return render(request, "accounting/ledger.html", {
        **data,
        "settings_obj": data["settings"],
        "account_type": account_type,
        "account_label": (
            "Cash at Hand"
            if account_type == CashTransaction.ACCOUNT_CASH
            else data["settings"].bank_label
        ),
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "today": timezone.localdate(),
    })


def _format_excel_amount(cell):
    cell.number_format = '#,##0.00'


def _download_workbook(workbook, filename):
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@role_required(*ROLES)
def import_template(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    headers = [
        "Type",
        "Date",
        "Account",
        "Category",
        "Party",
        "Amount",
        "Reference",
        "Description",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="0D6B57")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    widths = [14, 14, 16, 26, 30, 16, 22, 42]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    for row in range(2, 1002):
        sheet.cell(row=row, column=2).number_format = "yyyy-mm-dd"
        _format_excel_amount(sheet.cell(row=row, column=6))

    type_validation = DataValidation(
        type="list", formula1='"Receipt,Payment"', allow_blank=False
    )
    account_validation = DataValidation(
        type="list", formula1='"Cash,Bank"', allow_blank=False
    )
    sheet.add_data_validation(type_validation)
    sheet.add_data_validation(account_validation)
    type_validation.add("A2:A1001")
    account_validation.add("C2:C1001")

    guide = workbook.create_sheet("How to use")
    guide.append(["Simple cash book import"])
    guide["A1"].font = Font(size=16, bold=True, color="0D6B57")
    guide.append([])
    guide.append([
        "1. Enter one transaction on each row in the Transactions sheet.",
        "2. Use Receipt for money received and Payment for money paid.",
        "3. Type the date the money actually moved.",
        "4. Cash = Cash at Hand. Bank = Cash at Bank.",
        "5. Amount must be greater than zero.",
    ])
    guide.append([])
    guide.append(["Allowed categories"])
    guide["A7"].font = Font(bold=True)
    for key, label in CashTransaction.CATEGORY_CHOICES:
        guide.append([label, key])
    guide.column_dimensions["A"].width = 75
    guide.column_dimensions["B"].width = 24
    return _download_workbook(workbook, "accounting_import_template.xlsx")


def _cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _import_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, (int, float)):
        from openpyxl.utils.datetime import from_excel

        try:
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime.datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    raw = _cell_text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(raw, pattern).date()
        except ValueError:
            continue
    return None


def _import_money(value):
    raw = _cell_text(value).replace(",", "").replace("৳", "").replace("BDT", "").strip()
    try:
        amount = Decimal(raw)
    except (InvalidOperation, ValueError):
        return None
    if not amount.is_finite() or amount <= 0 or amount > Decimal("999999999999.99"):
        return None
    return amount.quantize(Decimal("0.01"))


def _normalise(value):
    return " ".join(_cell_text(value).lower().replace("&", "and").split())


TYPE_LOOKUP = {
    "receipt": CashTransaction.TYPE_RECEIPT,
    "money received": CashTransaction.TYPE_RECEIPT,
    "money in": CashTransaction.TYPE_RECEIPT,
    "payment": CashTransaction.TYPE_PAYMENT,
    "money paid": CashTransaction.TYPE_PAYMENT,
    "money out": CashTransaction.TYPE_PAYMENT,
}
ACCOUNT_LOOKUP = {
    "cash": CashTransaction.ACCOUNT_CASH,
    "cash at hand": CashTransaction.ACCOUNT_CASH,
    "hand": CashTransaction.ACCOUNT_CASH,
    "bank": CashTransaction.ACCOUNT_BANK,
    "cash at bank": CashTransaction.ACCOUNT_BANK,
    "bank account": CashTransaction.ACCOUNT_BANK,
}
CATEGORY_LOOKUP = {
    _normalise(key): key for key, _ in CashTransaction.CATEGORY_CHOICES
}
CATEGORY_LOOKUP.update({
    _normalise(label): key for key, label in CashTransaction.CATEGORY_CHOICES
})


def _run_import(school, upload, user):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
    except Exception:
        return {"fatal": "Could not read the file. Please upload a valid .xlsx file."}

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        workbook.close()
        return {"fatal": "The spreadsheet is empty."}
    columns = {
        _normalise(value).replace(" ", "_"): index
        for index, value in enumerate(header)
        if _cell_text(value)
    }
    required = ["type", "date", "account", "category", "party", "amount"]
    missing = [name for name in required if name not in columns]
    if missing:
        workbook.close()
        return {
            "fatal": "Missing column(s): " + ", ".join(name.title() for name in missing)
        }

    def value(row, name):
        index = columns.get(name)
        return row[index] if index is not None and index < len(row) else None

    settings_obj = get_settings(school)
    valid_rows = []
    errors = []
    for row_number, row in enumerate(rows, start=2):
        if row is None or all(_cell_text(value) == "" for value in row):
            continue
        if len(valid_rows) + len(errors) >= 2000:
            workbook.close()
            return {"fatal": "Too many rows. Import up to 2,000 transactions at a time."}
        row_type = TYPE_LOOKUP.get(_normalise(value(row, "type")))
        row_account = ACCOUNT_LOOKUP.get(_normalise(value(row, "account")))
        row_category = CATEGORY_LOOKUP.get(_normalise(value(row, "category")))
        row_date = _import_date(value(row, "date"))
        row_amount = _import_money(value(row, "amount"))
        party = _cell_text(value(row, "party"))[:180]
        row_errors = []
        if not row_type:
            row_errors.append("Type must be Receipt or Payment")
        if not row_date:
            row_errors.append("Date is not valid")
        elif row_date < settings_obj.opening_date:
            row_errors.append(
                f"Date is before the cash book start date ({settings_obj.opening_date:%d %b %Y})"
            )
        elif row_date > timezone.localdate():
            row_errors.append("Date cannot be in the future")
        if not row_account:
            row_errors.append("Account must be Cash or Bank")
        if not row_category:
            row_errors.append("Category is not recognised")
        if not party:
            row_errors.append("Party is required")
        if row_amount is None:
            row_errors.append("Amount must be greater than zero")
        if row_errors:
            errors.append((row_number, "; ".join(row_errors)))
            continue
        valid_rows.append({
            "transaction_type": row_type,
            "entry_date": row_date,
            "account_type": row_account,
            "category": row_category,
            "party": party,
            "amount": row_amount,
            "reference": _cell_text(value(row, "reference"))[:120],
            "description": _cell_text(value(row, "description")),
        })
    workbook.close()

    if not valid_rows and not errors:
        return {"fatal": "No transaction rows were found."}
    if errors:
        return {
            "fatal": "Nothing was imported because some rows need correction.",
            "errors": errors[:50],
            "error_total": len(errors),
        }
    with transaction.atomic():
        for row in valid_rows:
            create_transaction(
                school=school,
                created_by=user,
                **row,
            )
    return {"imported": len(valid_rows)}


@role_required(*ROLES)
def import_entries(request):
    result = None
    if request.method == "POST":
        upload = request.FILES.get("file")
        if not upload or not upload.name.lower().endswith(".xlsx"):
            messages.error(request, "Please choose an .xlsx Excel file.")
        elif upload.size > MAX_UPLOAD_BYTES:
            messages.error(request, "The file is too large. Maximum size is 5 MB.")
        else:
            result = _run_import(request.school, upload, request.user)
            if result.get("fatal"):
                messages.error(request, result["fatal"])
            else:
                messages.success(
                    request,
                    f"{result['imported']} transactions imported successfully.",
                )
                return redirect("accounting:entries")
    return render(request, "accounting/import.html", {
        "result": result,
        "categories": CashTransaction.CATEGORY_CHOICES,
    })


@role_required(*ROLES)
def export_entries(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    qs, _ = _filtered_entries(request)
    export_format = request.GET.get("format", "xlsx").lower()
    headers = [
        "Voucher No",
        "Type",
        "Date",
        "Account",
        "Category",
        "Party",
        "Amount",
        "Reference",
        "Description",
        "Source",
        "Status",
    ]
    if export_format == "csv":
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = (
            f'attachment; filename="accounting_transactions_{timezone.localdate():%Y%m%d}.csv"'
        )
        response.write("\ufeff")
        writer = csv.writer(response)
        writer.writerow(headers)
        for entry in qs.iterator():
            writer.writerow([
                entry.voucher_no,
                entry.get_transaction_type_display(),
                entry.entry_date,
                entry.account_label,
                entry.get_category_display(),
                entry.party,
                entry.amount,
                entry.reference,
                entry.description,
                entry.get_source_display(),
                "Voided" if entry.is_voided else "Active",
            ])
        return response

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="0D6B57")
        cell.font = Font(color="FFFFFF", bold=True)
    for entry in qs.iterator():
        sheet.append([
            entry.voucher_no,
            "Receipt" if entry.transaction_type == "receipt" else "Payment",
            entry.entry_date,
            entry.account_label,
            entry.get_category_display(),
            entry.party,
            entry.amount,
            entry.reference,
            entry.description,
            entry.get_source_display(),
            "Voided" if entry.is_voided else "Active",
        ])
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=3).number_format = "yyyy-mm-dd"
        _format_excel_amount(sheet.cell(row=row, column=7))
    widths = [14, 12, 14, 18, 26, 28, 16, 22, 42, 22, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    return _download_workbook(
        workbook,
        f"accounting_transactions_{timezone.localdate():%Y%m%d}.xlsx",
    )
