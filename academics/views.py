import datetime
import re
from io import BytesIO

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.hashers import make_password
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import ProtectedError, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from accounts.decorators import role_required
from .models import AcademicYear, SchoolClass, Student

User = get_user_model()


def default_password():
    return getattr(settings, "DEFAULT_STUDENT_PASSWORD", "12345678")


def _to_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _limit_message(school):
    return (
        f"Your {school.effective_plan_label} plan allows up to {school.student_limit} students "
        f"and you already have {school.students_used()}. Contact us to upgrade your plan."
    )


# ---------------------------------------------------------------------------
# Academic Years
# ---------------------------------------------------------------------------

def _year_clean(request, school, exclude_pk=None):
    name = request.POST.get("name", "").strip()
    start = request.POST.get("start_date", "").strip()
    end = request.POST.get("end_date", "").strip()
    errors = []
    start_d = end_d = None

    if not name:
        errors.append("Year name is required.")
    elif AcademicYear.objects.filter(school=school, name=name).exclude(pk=exclude_pk).exists():
        errors.append(f"Academic year '{name}' already exists.")

    try:
        start_d = datetime.date.fromisoformat(start) if start else None
        end_d = datetime.date.fromisoformat(end) if end else None
    except ValueError:
        errors.append("Invalid date.")

    if start_d and end_d and end_d < start_d:
        errors.append("End date must be after the start date.")

    return {"name": name, "start_date": start_d, "end_date": end_d}, errors


@role_required("admin")
def year_list(request):
    school = request.school
    errors = []
    form = {}

    if request.method == "POST":
        form = request.POST
        data, errors = _year_clean(request, school)
        if not errors:
            with transaction.atomic():
                make_current = (
                    request.POST.get("is_current") == "on"
                    or not AcademicYear.objects.filter(school=school).exists()
                )
                if make_current:
                    AcademicYear.objects.filter(school=school).update(is_current=False)
                AcademicYear.objects.create(school=school, is_current=make_current, **data)
            messages.success(request, f"Academic year '{data['name']}' added.")
            return redirect("academics:year_list")

    years = AcademicYear.objects.filter(school=school)
    return render(request, "academics/years.html", {
        "years": years, "errors": errors, "form": form,
    })


@role_required("admin")
def year_edit(request, pk):
    school = request.school
    year = get_object_or_404(AcademicYear, pk=pk, school=school)
    errors = []

    if request.method == "POST":
        data, errors = _year_clean(request, school, exclude_pk=year.pk)
        if not errors:
            with transaction.atomic():
                year.name = data["name"]
                year.start_date = data["start_date"]
                year.end_date = data["end_date"]
                if request.POST.get("is_current") == "on":
                    AcademicYear.objects.filter(school=school).update(is_current=False)
                    year.is_current = True
                year.save()
            messages.success(request, "Academic year updated.")
            return redirect("academics:year_list")
        form = request.POST
    else:
        form = {
            "name": year.name,
            "start_date": year.start_date.isoformat() if year.start_date else "",
            "end_date": year.end_date.isoformat() if year.end_date else "",
            "is_current": "on" if year.is_current else "",
        }

    return render(request, "academics/year_edit.html", {
        "year": year, "errors": errors, "form": form,
    })


@role_required("admin")
@require_POST
def year_delete(request, pk):
    year = get_object_or_404(AcademicYear, pk=pk, school=request.school)
    try:
        year.delete()
        messages.success(request, f"Academic year '{year.name}' deleted.")
    except ProtectedError:
        messages.error(request, "This year is in use (students or fees are linked to it) and cannot be deleted.")
    return redirect("academics:year_list")


@role_required("admin")
@require_POST
def year_current(request, pk):
    school = request.school
    year = get_object_or_404(AcademicYear, pk=pk, school=school)
    with transaction.atomic():
        AcademicYear.objects.filter(school=school).update(is_current=False)
        year.is_current = True
        year.save()
    messages.success(request, f"'{year.name}' is now the current academic year.")
    return redirect("academics:year_list")


# ---------------------------------------------------------------------------
# Classes
# ---------------------------------------------------------------------------

def _class_clean(request, school, exclude_pk=None):
    name = request.POST.get("name", "").strip()
    section = request.POST.get("section", "").strip()
    order_raw = request.POST.get("order", "").strip()
    errors = []

    if not name:
        errors.append("Class name is required.")
    elif SchoolClass.objects.filter(school=school, name=name, section=section).exclude(pk=exclude_pk).exists():
        errors.append("This class and section already exist.")

    order = 0
    if order_raw:
        order = _to_int(order_raw)
        if order is None or order < 0:
            errors.append("Order must be a positive number.")
            order = 0

    return {"name": name, "section": section, "order": order}, errors


@role_required("admin")
def class_list(request):
    school = request.school
    errors = []
    form = {}

    if request.method == "POST":
        form = request.POST
        data, errors = _class_clean(request, school)
        if not errors:
            SchoolClass.objects.create(school=school, **data)
            messages.success(request, f"Class '{data['name']}' added.")
            return redirect("academics:class_list")

    classes = SchoolClass.objects.filter(school=school)
    return render(request, "academics/classes.html", {
        "classes": classes, "errors": errors, "form": form,
    })


@role_required("admin")
def class_edit(request, pk):
    school = request.school
    obj = get_object_or_404(SchoolClass, pk=pk, school=school)
    errors = []

    if request.method == "POST":
        data, errors = _class_clean(request, school, exclude_pk=obj.pk)
        if not errors:
            obj.name = data["name"]
            obj.section = data["section"]
            obj.order = data["order"]
            obj.save()
            messages.success(request, "Class updated.")
            return redirect("academics:class_list")
        form = request.POST
    else:
        form = {"name": obj.name, "section": obj.section, "order": obj.order}

    return render(request, "academics/class_edit.html", {
        "obj": obj, "errors": errors, "form": form,
    })


@role_required("admin")
@require_POST
def class_delete(request, pk):
    obj = get_object_or_404(SchoolClass, pk=pk, school=request.school)
    try:
        obj.delete()
        messages.success(request, f"Class '{obj.label}' deleted.")
    except ProtectedError:
        messages.error(request, "This class has students or fees and cannot be deleted.")
    return redirect("academics:class_list")


# ---------------------------------------------------------------------------
# Students
# ---------------------------------------------------------------------------

def _student_clean(request, school, instance=None):
    p = request.POST
    errors = []

    student_id = instance.student_id if instance else p.get("student_id", "").strip()
    name = p.get("name", "").strip()
    father_name = p.get("father_name", "").strip()
    phone = p.get("phone", "").strip()
    roll_raw = p.get("roll", "").strip()

    school_class = SchoolClass.objects.filter(
        school=school, pk=_to_int(p.get("school_class"))
    ).first()
    academic_year = AcademicYear.objects.filter(
        school=school, pk=_to_int(p.get("academic_year"))
    ).first()

    if not student_id:
        errors.append("Student ID is required.")
    elif instance is None:
        if Student.objects.filter(school=school, student_id=student_id).exists():
            errors.append(f"Student ID '{student_id}' already exists.")
        elif User.objects.filter(school=school, username=student_id).exists():
            errors.append(f"'{student_id}' is already used as a login username.")

    if not name:
        errors.append("Student name is required.")
    if school_class is None:
        errors.append("Please select a class.")
    if academic_year is None:
        errors.append("Please select an academic year.")

    roll = None
    if roll_raw:
        roll = _to_int(roll_raw)
        if roll is None or roll < 0:
            errors.append("Roll must be a number.")
            roll = None

    data = {
        "student_id": student_id, "name": name, "father_name": father_name,
        "phone": phone, "roll": roll,
        "school_class": school_class, "academic_year": academic_year,
    }
    return data, errors


def _student_form_context(school, form, errors, student=None):
    return {
        "student": student, "form": form, "errors": errors,
        "classes": SchoolClass.objects.filter(school=school),
        "years": AcademicYear.objects.filter(school=school),
    }


@role_required("admin", "accountant", "teacher")
def student_list(request):
    school = request.school
    qs = Student.objects.filter(school=school).select_related("school_class", "academic_year")

    q = request.GET.get("q", "").strip()
    class_id = request.GET.get("class", "")
    year_id = request.GET.get("year", "")

    if q:
        qs = qs.filter(
            Q(name__icontains=q) | Q(student_id__icontains=q)
            | Q(father_name__icontains=q) | Q(phone__icontains=q)
        )
    if class_id.isdigit():
        qs = qs.filter(school_class_id=int(class_id))
    if year_id.isdigit():
        qs = qs.filter(academic_year_id=int(year_id))

    page_obj = Paginator(qs, 50).get_page(request.GET.get("page"))

    params = request.GET.copy()
    params.pop("page", None)

    return render(request, "academics/students.html", {
        "page_obj": page_obj,
        "q": q, "class_id": class_id, "year_id": year_id,
        "classes": SchoolClass.objects.filter(school=school),
        "years": AcademicYear.objects.filter(school=school),
        "querystring": params.urlencode(),
    })


@role_required("admin")
def student_add(request):
    school = request.school
    errors = []

    # Plan limit: stop before the admin fills in the whole form.
    if not school.can_add_students(1):
        messages.error(request, _limit_message(school))
        return redirect("academics:student_list")

    if request.method == "POST":
        form = request.POST
        data, errors = _student_clean(request, school)
        if not errors:
            with transaction.atomic():
                user = User(
                    school=school, username=data["student_id"],
                    role="student", first_name=data["name"][:150],
                )
                user.password = make_password(default_password())
                user.save()
                Student.objects.create(school=school, user=user, **data)
            messages.success(
                request,
                f"Student added. Login: {data['student_id']} / password: {default_password()}",
            )
            return redirect("academics:student_list")
    else:
        current = AcademicYear.objects.filter(school=school, is_current=True).first()
        form = {"academic_year": str(current.pk) if current else ""}

    return render(request, "academics/student_form.html",
                  _student_form_context(school, form, errors))


@role_required("admin")
def student_edit(request, pk):
    school = request.school
    student = get_object_or_404(Student, pk=pk, school=school)
    errors = []

    if request.method == "POST":
        form = request.POST
        data, errors = _student_clean(request, school, instance=student)
        if not errors:
            is_active = request.POST.get("is_active") == "on"
            with transaction.atomic():
                student.name = data["name"]
                student.father_name = data["father_name"]
                student.phone = data["phone"]
                student.roll = data["roll"]
                student.school_class = data["school_class"]
                student.academic_year = data["academic_year"]
                student.is_active = is_active
                student.save()
                if student.user:
                    student.user.first_name = data["name"][:150]
                    student.user.is_active = is_active
                    student.user.save(update_fields=["first_name", "is_active"])
            messages.success(request, "Student updated.")
            return redirect("academics:student_list")
    else:
        form = {
            "student_id": student.student_id, "name": student.name,
            "father_name": student.father_name, "phone": student.phone,
            "roll": student.roll if student.roll is not None else "",
            "school_class": str(student.school_class_id),
            "academic_year": str(student.academic_year_id),
            "is_active": "on" if student.is_active else "",
        }

    return render(request, "academics/student_form.html",
                  _student_form_context(school, form, errors, student=student))


@role_required("admin")
@require_POST
def student_delete(request, pk):
    student = get_object_or_404(Student, pk=pk, school=request.school)
    user = student.user
    try:
        with transaction.atomic():
            student.delete()
            if user:
                user.delete()
        messages.success(request, f"Student '{student.name}' deleted.")
    except ProtectedError:
        messages.error(request, "This student has payment records and cannot be deleted.")
    return redirect("academics:student_list")


@role_required("admin")
@require_POST
def student_reset_password(request, pk):
    student = get_object_or_404(Student, pk=pk, school=request.school)
    if student.user:
        student.user.set_password(default_password())
        student.user.save(update_fields=["password"])
        messages.success(
            request,
            f"Password for {student.student_id} reset to: {default_password()}",
        )
    else:
        messages.error(request, "This student has no login account.")
    return redirect("academics:student_list")


# ---------------------------------------------------------------------------
# Excel import
# ---------------------------------------------------------------------------

REQUIRED_COLS = ["student_id", "name", "class"]
MAX_UPLOAD_BYTES = 5 * 1024 * 1024


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _class_order(name):
    match = re.search(r"\d+", name)
    return int(match.group()) if match else 0


def _run_import(school, year, upload, limit=None):
    """
    limit = how many more students the school's plan allows (None = unlimited).
    Rows beyond the limit are skipped and listed in the result.
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(upload, read_only=True, data_only=True)
    except Exception:
        return {"fatal": "Could not read the file. Please upload a valid .xlsx file."}

    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        wb.close()
        return {"fatal": "The file is empty."}

    cols = {}
    for i, h in enumerate(header):
        key = _cell(h).lower().replace(" ", "_")
        if key:
            cols[key] = i

    missing = [c for c in REQUIRED_COLS if c not in cols]
    if missing:
        wb.close()
        return {"fatal": "Missing column(s): " + ", ".join(missing)}

    def get(row, key):
        i = cols.get(key)
        return _cell(row[i]) if i is not None and i < len(row) else ""

    existing_ids = set(Student.objects.filter(school=school).values_list("student_id", flat=True))
    existing_users = set(User.objects.filter(school=school).values_list("username", flat=True))

    skipped = []
    to_create = []
    seen = set()

    for row_no, row in enumerate(rows, start=2):
        if row is None or all(_cell(v) == "" for v in row):
            continue

        sid = get(row, "student_id")
        name = get(row, "name")
        cname = get(row, "class")
        section = get(row, "section")
        roll_raw = get(row, "roll")

        if not sid or not name or not cname:
            skipped.append((row_no, "student_id, name and class are required"))
            continue
        if sid in existing_ids or sid in existing_users or sid in seen:
            skipped.append((row_no, f"Student ID {sid} already exists"))
            continue

        roll = None
        if roll_raw:
            if not roll_raw.isdigit():
                skipped.append((row_no, "Roll must be a number"))
                continue
            roll = int(roll_raw)

        seen.add(sid)
        to_create.append({
            "row": row_no,
            "student_id": sid, "name": name, "class": cname, "section": section,
            "father_name": get(row, "father_name"), "phone": get(row, "phone"),
            "roll": roll,
        })

    wb.close()

    # Plan limit: keep only as many valid rows as the plan still allows.
    limit_skipped = 0
    if limit is not None and len(to_create) > limit:
        for item in to_create[limit:]:
            skipped.append((item["row"], "Student limit of your plan reached"))
        limit_skipped = len(to_create) - limit
        to_create = to_create[:limit]

    new_classes = 0
    if to_create:
        with transaction.atomic():
            class_cache = {
                (c.name.lower(), c.section.lower()): c
                for c in SchoolClass.objects.filter(school=school)
            }
            for item in to_create:
                key = (item["class"].lower(), item["section"].lower())
                if key not in class_cache:
                    class_cache[key] = SchoolClass.objects.create(
                        school=school, name=item["class"], section=item["section"],
                        order=_class_order(item["class"]),
                    )
                    new_classes += 1

            # Hash the default password once, not once per student (much faster).
            hashed = make_password(default_password())
            User.objects.bulk_create([
                User(school=school, username=i["student_id"], role="student",
                     first_name=i["name"][:150], password=hashed)
                for i in to_create
            ], batch_size=500)

            user_ids = dict(
                User.objects.filter(school=school, role="student")
                .values_list("username", "pk")
            )

            Student.objects.bulk_create([
                Student(
                    school=school, user_id=user_ids.get(i["student_id"]),
                    student_id=i["student_id"], name=i["name"],
                    father_name=i["father_name"], phone=i["phone"], roll=i["roll"],
                    school_class=class_cache[(i["class"].lower(), i["section"].lower())],
                    academic_year=year,
                )
                for i in to_create
            ], batch_size=500)

    skipped.sort(key=lambda item: item[0])
    return {
        "created": len(to_create),
        "new_classes": new_classes,
        "skipped": skipped[:200],
        "skipped_total": len(skipped),
        "limit_skipped": limit_skipped,
    }


@role_required("admin")
def student_import(request):
    school = request.school
    years = AcademicYear.objects.filter(school=school)
    result = None

    if request.method == "POST":
        upload = request.FILES.get("file")
        year = years.filter(pk=_to_int(request.POST.get("academic_year"))).first()
        remaining = school.students_remaining()

        if year is None:
            messages.error(request, "Please select an academic year.")
        elif not upload or not upload.name.lower().endswith(".xlsx"):
            messages.error(request, "Please choose an .xlsx Excel file.")
        elif upload.size > MAX_UPLOAD_BYTES:
            messages.error(request, "File is too large (max 5 MB).")
        elif remaining == 0:
            messages.error(request, _limit_message(school))
        else:
            result = _run_import(school, year, upload, limit=remaining)
            if result.get("fatal"):
                messages.error(request, result["fatal"])
                result = None

    return render(request, "academics/student_import.html", {
        "years": years, "result": result,
        "default_password": default_password(),
        "quota": school.plan_info,  # read after the import so the numbers are fresh
    })


@role_required("admin")
def student_import_sample(request):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(["student_id", "name", "father_name", "phone", "class", "section", "roll"])
    ws.append(["1001", "Rahim Uddin", "Abdul Karim", "01711000000", "Class 6", "A", 1])
    ws.append(["1002", "Fatema Akter", "Md. Jamal", "01811000000", "Class 6", "A", 2])

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="students_sample.xlsx"'
    return response
